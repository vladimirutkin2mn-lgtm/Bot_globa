import csv,json,re,requests,sys,time
from pathlib import Path
from urllib.parse import urljoin
from bs4 import BeautifulSoup
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent; OUT=ROOT/'output_gisp_1000'; OUT.mkdir(exist_ok=True)
inns=[]
for p in sorted(ROOT.glob('inns_1000.part*')): inns += [x.strip() for x in p.read_text().splitlines() if x.strip()]
T=set(inns); assert len(inns)==1000
S=requests.Session(); S.headers.update({'User-Agent':'Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/140 Safari/537.36','Accept':'*/*'})
BASES=['https://portfolio.gisp.gov.ru/pp719v2/pub/prod/','https://gisp.gov.ru/pp719v2/pub/prod/']
diag={'pages':[],'candidates':[],'download':None}
cands=[]
def add(url,label=''):
    if not url: return
    url=urljoin(current_base,url)
    if url not in [x[0] for x in cands]: cands.append((url,label))
for base in BASES:
    current_base=base
    try:
        r=S.get(base,timeout=30,allow_redirects=True); diag['pages'].append({'url':base,'status':r.status_code,'len':len(r.text),'final':r.url})
        soup=BeautifulSoup(r.text,'html.parser')
        for a in soup.find_all(['a','button'],href=True):
            txt=' '.join(a.stripped_strings).lower(); href=a.get('href','')
            if any(x in (txt+' '+href.lower()) for x in ['xlsx','xls','скач','выгруз','excel','export']): add(href,txt)
        for tag in soup.find_all(attrs={'onclick':True}):
            txt=' '.join(tag.stripped_strings).lower(); oc=tag.get('onclick','')
            for q in re.findall(r"['\"]([^'\"]+)['\"]",oc):
                if any(x in (txt+' '+q.lower()) for x in ['xlsx','xls','скач','выгруз','excel','export']): add(q,txt)
        for m in re.findall(r"(?:https?://[^\"'<>\s]+|/[^\"'<>\s]+)(?:xlsx|xls|excel|export|download)[^\"'<>\s]*",r.text,re.I): add(m,'regex')
        # inspect referenced JS for export endpoints
        for sc in soup.find_all('script',src=True)[:30]:
            try:
                js=S.get(urljoin(base,sc['src']),timeout=15).text
                for m in re.findall(r"['\"]([^'\"]*(?:xlsx|excel|export|download)[^'\"]*)['\"]",js,re.I): add(m,'js')
            except: pass
    except Exception as e: diag['pages'].append({'url':base,'error':repr(e)})
# prioritize active/current download labels
cands.sort(key=lambda x:(0 if ('действ' in x[1] or 'active' in x[1]) else 1,0 if 'xlsx' in x[0].lower() else 1))
# common fallbacks
for base in BASES:
    current_base=base
    for q in ['?download=xlsx&active=1','?export=xlsx&active=1','export/?format=xlsx&active=1','export.xlsx','download.xlsx','excel/','export/']:
        add(q,'fallback')
diag['candidates']=[{'url':u,'label':l} for u,l in cands[:100]]
path=OUT/'gisp.xlsx'
for u,lbl in cands[:100]:
    try:
        rr=S.get(u,timeout=90,allow_redirects=True,stream=True)
        first=next(rr.iter_content(4096),b'')
        ct=rr.headers.get('content-type','').lower()
        if rr.status_code==200 and (first.startswith(b'PK') or 'spreadsheet' in ct or 'excel' in ct):
            with path.open('wb') as f:
                f.write(first)
                for chunk in rr.iter_content(1024*1024): f.write(chunk)
            if path.stat().st_size>10000:
                diag['download']={'url':u,'label':lbl,'size':path.stat().st_size,'content_type':ct}; break
    except Exception as e: pass
res={i:{'ИНН':i,'ГИСП статус источника':'ERROR_NO_XLSX','ГИСП найден':'','ГИСП записей':0,'ГИСП продукция':'','ГИСП источник':''} for i in inns}
if path.exists() and path.stat().st_size>10000:
    try:
        wb=load_workbook(path,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]
        header_row=None; inn_col=None; prod_cols=[]; name_cols=[]
        for ridx,row in enumerate(ws.iter_rows(min_row=1,max_row=15,values_only=True),1):
            vals=[str(v).strip() if v is not None else '' for v in row]
            for j,v in enumerate(vals):
                vl=v.lower().replace(' ','')
                if 'инн'==vl or ('инн' in vl and len(vl)<30): inn_col=j
            if inn_col is not None:
                header_row=ridx
                for j,v in enumerate(vals):
                    vl=v.lower()
                    if 'продукц' in vl or 'наименование продукции' in vl: prod_cols.append(j)
                    if 'наименование' in vl and j!=inn_col: name_cols.append(j)
                break
        if inn_col is None: raise RuntimeError('INN column not found')
        agg={}
        for row in ws.iter_rows(min_row=header_row+1,values_only=True):
            if inn_col>=len(row): continue
            inn=re.sub(r'\D','',str(row[inn_col] or ''))
            if inn not in T: continue
            a=agg.setdefault(inn,{'n':0,'products':[]}) ; a['n']+=1
            vals=[]
            for j in prod_cols:
                if j<len(row) and row[j]: vals.append(str(row[j]).strip())
            if not vals:
                for j in name_cols[-2:]:
                    if j<len(row) and row[j]: vals.append(str(row[j]).strip())
            for v in vals:
                if v and v not in a['products'] and len(a['products'])<12: a['products'].append(v)
        for inn in inns:
            if inn in agg:
                res[inn].update({'ГИСП статус источника':'OK','ГИСП найден':'YES','ГИСП записей':agg[inn]['n'],'ГИСП продукция':' | '.join(agg[inn]['products']),'ГИСП источник':diag['download']['url']})
            else:
                res[inn].update({'ГИСП статус источника':'OK','ГИСП найден':'NO','ГИСП источник':diag['download']['url']})
        diag['matched']=len(agg); diag['header_row']=header_row; diag['inn_col']=inn_col+1; diag['prod_cols']=[x+1 for x in prod_cols]
    except Exception as e:
        diag['parse_error']=repr(e)
        for inn in inns: res[inn]['ГИСП статус источника']='ERROR_PARSE'
fields=list(next(iter(res.values())).keys())
with (OUT/'gisp_1000.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=fields,delimiter=';'); w.writeheader(); w.writerows(res[i] for i in inns)
summary={'rows':1000,'download':diag.get('download'),'matched':sum(1 for i in inns if res[i]['ГИСП найден']=='YES'),'source_ok':sum(1 for i in inns if res[i]['ГИСП статус источника']=='OK'),'diagnostics':diag}
(OUT/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8'); print(json.dumps(summary,ensure_ascii=False,indent=2)[:20000])
